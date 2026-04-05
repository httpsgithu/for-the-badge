#!/usr/bin/env python3
"""
Docling processor for document text extraction
Supports: PDF, DOCX, BMP, GIF, JPEG, JPG, PNG, TIF, TIFF, WebP, and other formats
"""

import sys
import json
import os
import tempfile
from pathlib import Path

def extract_text_with_docling(file_path):
    """
    Extract text from document using Docling library
    """
    try:
        # Import Docling (will need to be installed: pip install docling)
        from docling.document_converter import DocumentConverter
        
        # Initialize converter
        converter = DocumentConverter()
        
        # Convert document
        result = converter.convert(file_path)
        
        # Extract text content
        if hasattr(result, 'document') and hasattr(result.document, 'export_to_text'):
            text = result.document.export_to_text()
        elif hasattr(result, 'text'):
            text = result.text
        elif hasattr(result, 'content'):
            text = result.content
        else:
            # Fallback - convert result to string
            text = str(result)
        
        return {
            "success": True,
            "text": text.strip() if text else "",
            "pages": getattr(result, 'page_count', 1),
            "format": Path(file_path).suffix.lower()
        }
        
    except ImportError:
        return {
            "success": False,
            "error": "Docling library not installed. Please install with: pip install docling"
        }
    except Exception as e:
        return {
            "success": False,
            "error": f"Docling processing failed: {str(e)}"
        }

def fallback_text_extraction(file_path):
    """
    Fallback text extraction using alternative libraries
    """
    file_extension = Path(file_path).suffix.lower()
    
    try:
        # PDF extraction using PyPDF2 or similar
        if file_extension == '.pdf':
            try:
                import PyPDF2
                with open(file_path, 'rb') as file:
                    reader = PyPDF2.PdfReader(file)
                    text = ""
                    for page in reader.pages:
                        text += page.extract_text()
                    return {
                        "success": True,
                        "text": text.strip(),
                        "pages": len(reader.pages),
                        "format": file_extension,
                        "method": "PyPDF2 fallback"
                    }
            except ImportError:
                pass
        
        # Image OCR using pytesseract
        elif file_extension in ['.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.tif', '.webp', '.gif']:
            try:
                import pytesseract
                from PIL import Image
                
                # Load and process image
                image = Image.open(file_path)
                text = pytesseract.image_to_string(image)
                
                return {
                    "success": True,
                    "text": text.strip(),
                    "pages": 1,
                    "format": file_extension,
                    "method": "Tesseract OCR fallback"
                }
            except ImportError:
                return {
                    "success": False,
                    "error": "Neither Docling nor pytesseract available for image processing"
                }
        
        # DOCX extraction using python-docx
        elif file_extension == '.docx':
            try:
                from docx import Document
                doc = Document(file_path)
                text = ""
                for paragraph in doc.paragraphs:
                    text += paragraph.text + "\n"
                
                return {
                    "success": True,
                    "text": text.strip(),
                    "pages": max(1, len(doc.paragraphs) // 20),  # Rough estimate
                    "format": file_extension,
                    "method": "python-docx fallback"
                }
            except ImportError:
                pass
        
        # Plain text files
        elif file_extension in ['.txt', '.md']:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
                return {
                    "success": True,
                    "text": text.strip(),
                    "pages": max(1, len(text) // 2000),  # Rough page estimate
                    "format": file_extension,
                    "method": "direct text read"
                }
        
        return {
            "success": False,
            "error": f"No fallback method available for format: {file_extension}"
        }
        
    except Exception as e:
        return {
            "success": False,
            "error": f"Fallback extraction failed: {str(e)}"
        }

def main():
    if len(sys.argv) != 2:
        result = {
            "success": False,
            "error": "Usage: python docling_processor.py <file_path>"
        }
        print(json.dumps(result))
        sys.exit(1)
    
    file_path = sys.argv[1]
    
    # Validate file exists
    if not os.path.exists(file_path):
        result = {
            "success": False,
            "error": f"File not found: {file_path}"
        }
        print(json.dumps(result))
        sys.exit(1)
    
    # Get file info
    file_size = os.path.getsize(file_path)
    file_extension = Path(file_path).suffix.lower()
    
    # Check file size (limit to 50MB for processing)
    if file_size > 50 * 1024 * 1024:
        result = {
            "success": False,
            "error": f"File too large: {file_size} bytes (max 50MB)"
        }
        print(json.dumps(result))
        sys.exit(1)
    
    # Supported formats
    supported_formats = [
        '.pdf', '.docx', '.jpg', '.jpeg', '.png', '.bmp', 
        '.tiff', '.tif', '.webp', '.gif', '.txt', '.md',
        '.xlsx', '.pptx', '.html', '.csv'
    ]
    
    if file_extension not in supported_formats:
        result = {
            "success": False,
            "error": f"Unsupported file format: {file_extension}. Supported formats: {', '.join(supported_formats)}"
        }
        print(json.dumps(result))
        sys.exit(1)
    
    # Try Docling first
    result = extract_text_with_docling(file_path)
    
    # If Docling fails, try fallback methods
    if not result["success"]:
        result = fallback_text_extraction(file_path)
    
    # Output result as JSON
    print(json.dumps(result))

if __name__ == "__main__":
    main()